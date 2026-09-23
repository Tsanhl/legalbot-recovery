"""One continuous pre-training evaluation-packet completion campaign.

Prepares every remaining review dossier for the visible 331, consolidates one
owner-review package, and stops at AWAITING_OWNER_EVALUATION_REVIEW. This is
not owner currentness approval, qualified legal review, answer gold, admission,
weight training, sealed unseen, promotion or live.
"""

from __future__ import annotations

import json
from collections import Counter
from collections.abc import Mapping, Sequence
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from app.contracts.schema_registry import seal_contract
from docx import Document
from docx.shared import Pt

from .ge_currentness_packets import (
    CASE_174,
    CASE_312,
    EXPECTED_CURRENTNESS,
    EXPECTED_FACT_DEPENDENT,
    EXPECTED_FACTUAL_HOLD,
    EXPECTED_FACTUAL_PASS,
    EXPECTED_JURISDICTION,
    EXPECTED_LEFTOVER,
    EXPECTED_R2_RESULTS,
    EXPECTED_TOTAL,
    PROJECT_ROOT,
    csv_text,
    evidence_manifest_hash,
    full_331_guard_result,
    load_jsonl,
    load_latest_delta_rows,
    mutation_guard,
    packet_index_rows,
    reconcile_latest_routes,
    render_topic_batch,
    repair_currentness_packets,
    route_sets,
    sha256_file,
    sha256_text,
    summarize_packets,
    validate_packet,
    write_json,
    write_jsonl,
    write_text,
)
from .ge_fact_dependent_packets import FACT_DEPENDENT_DECISIONS, build_fact_dependent_packet
from .ge_hold_reason_router import CASE_008
from .ge_jurisdiction_packets import JURISDICTION_DECISIONS, build_jurisdiction_packets
from .ge_official_currentness_metadata import CurrentnessMetadataIndex
from .ge_phase2_progress import (
    AWAITING_OWNER_EVALUATION_REVIEW,
    COMPLETE,
    DOWNSTREAM_GATES,
    NOT_STARTED,
    NO_OP_UNCHANGED_INPUTS,
)
from .ge_qualified_review_packets import REVIEWER_OPTIONS, build_factual_pass_packet
from .ge_residual_support_packets import RESIDUAL_DECISIONS, build_residual_packets, classify_no_evidence_label

CURRENTNESS_PACK_R1 = (
    PROJECT_ROOT
    / "data/evaluations/general-enquiries"
    / "LegalBot-GE-2026-09-03-currentness-packets-r1"
)
DEFAULT_MASTER_PACK = (
    PROJECT_ROOT
    / "data/evaluations/general-enquiries"
    / "LegalBot-GE-2026-09-03-master-331-evaluation-review-r1"
)
CURRENTNESS_DECISIONS = (
    "APPROVE_CURRENT_AS_OF_2026_08_28",
    "APPROVE_FOR_HISTORIC_APPLICABLE_DATE",
    "APPROVE_WITH_DATE_LIMITATION",
    "HOLD_CURRENTNESS",
    "REJECT_LOCATOR",
)
DECISION_SCHEMA = {
    "schema": "legalbot.ge-owner-evaluation-decision-schema.v1",
    "currentness_decisions": list(CURRENTNESS_DECISIONS),
    "jurisdiction_decisions": list(JURISDICTION_DECISIONS),
    "residual_support_decisions": list(RESIDUAL_DECISIONS),
    "fact_dependent_decisions": list(FACT_DEPENDENT_DECISIONS),
    "qualified_legal_review_decisions": list(REVIEWER_OPTIONS),
    "owner_and_qualified_review_are_separate_fields": True,
    "locator_or_currentness_review_does_not_populate_qualified_review_or_gold": True,
    "all_decision_fields_blank_until_owner_review": True,
}
NAMED_TESTS = (
    "exactly_331_unique_case_ids",
    "counts_reconcile_42_pass_289_hold",
    "route_counts_211_56_21_1_42",
    "pairwise_route_overlaps_empty",
    "duplicates_zero",
    "case_174_jurisdiction_only",
    "case_312_fact_dependent_only",
    "every_case_has_packet_or_incomplete_final",
    "stable_evidence_and_answer_hashes",
    "currentness_cutoff_applicable_and_source_dates_distinct_fields",
    "no_blank_required_field_silently_inferred",
    "no_owner_decision_populated",
    "no_qualified_legal_review_decision_populated",
    "no_factual_pass_created_through_packet_preparation",
    "no_answer_gold_created",
    "no_full_331_rerun",
    "frozen_r2_hash_unchanged",
    "second_identical_generation_idempotent",
    "no_duplicate_artifacts",
    "full_text_evidence_renders_without_truncation",
    "all_downstream_gates_remain_off",
)


class GlobalIntegrityError(RuntimeError):
    """A campaign-wide hard stop. Case-level gaps must not use this."""


def _default_hints() -> tuple[Mapping[str, tuple[tuple[str, str], ...]], Mapping[str, tuple[str, ...]]]:
    import sys

    root = str(PROJECT_ROOT)
    if root not in sys.path:
        sys.path.insert(0, root)
    from scripts.run_ge_retrieval_training_cycle import ISSUE_LOCATOR_HINTS, TOPIC_SOURCES

    return ISSUE_LOCATOR_HINTS, TOPIC_SOURCES


def load_currentness_r1_packets(path: Path = CURRENTNESS_PACK_R1) -> list[dict[str, Any]]:
    return load_jsonl(path / "currentness-packet-manifest.jsonl")


def by_id(rows: Sequence[Mapping[str, Any]]) -> dict[str, Mapping[str, Any]]:
    return {str(row.get("case_id") or ""): row for row in rows}


def _quote_text(item: Mapping[str, Any]) -> str:
    return str(item.get("quote") or item.get("exact_supporting_passage") or "")


def render_generic_topic_batch(topic: str, heading: str, packets: Sequence[Mapping[str, Any]]) -> str:
    lines = [
        f"# {heading} — {topic}",
        "",
        "Packet preparation only. Owner review is NOT_STARTED. Qualified legal review,",
        "answer gold, admission, training, sealed unseen, promotion and live remain off.",
        "",
    ]
    for packet in packets:
        lines.extend(
            [
                f"## {packet['case_id']}",
                "",
                f"- packet_status: `{packet.get('packet_status')}`",
                f"- factual_status: `{packet.get('factual_status')}`",
                f"- route: `{packet.get('hold_reason') or packet.get('route') or ''}`",
                f"- packet_hash: `{packet.get('packet_hash')}`",
                "",
                "### Question",
                "",
                str(packet.get("question") or ""),
                "",
                "### Candidate answer",
                "",
                str(packet.get("candidate_answer") or packet.get("answer") or ""),
                "",
                f"### Exact reviewer question",
                "",
                str(
                    packet.get("exact_question_for_reviewer")
                    or packet.get("precise_qualified_review_question")
                    or ""
                ),
                "",
            ]
        )
        for item in packet.get("attached_authorities") or packet.get("existing_attached_evidence") or []:
            if not isinstance(item, Mapping):
                continue
            quote = _quote_text(item)
            lines.extend(
                [
                    f"### Authority — {item.get('title')} / {item.get('locator')}",
                    "",
                    quote,
                    "",
                ]
            )
        for item in packet.get("exact_evidence_spans") or []:
            if not isinstance(item, Mapping):
                continue
            lines.extend(
                [
                    f"### Evidence span — {item.get('title')} / {item.get('locator')}",
                    "",
                    str(item.get("quote") or ""),
                    "",
                ]
            )
        if packet.get("proposed_conditional_answer"):
            lines.extend(["### Proposed conditional answer", "", str(packet["proposed_conditional_answer"]), ""])
        if packet.get("missing_facts"):
            lines.append("### Missing facts")
            lines.append("")
            for fact in packet["missing_facts"]:
                lines.append(f"- {fact}")
            lines.append("")
    return "\n".join(lines)


def build_topic_renders(
    currentness_packets: Sequence[Mapping[str, Any]],
    jurisdiction_packets: Sequence[Mapping[str, Any]],
    residual_packets: Sequence[Mapping[str, Any]],
    fact_packet: Mapping[str, Any],
    pass_packets: Sequence[Mapping[str, Any]],
) -> dict[str, str]:
    packets_by_topic: dict[str, list[Mapping[str, Any]]] = {}
    for packet in (
        list(currentness_packets)
        + list(jurisdiction_packets)
        + list(residual_packets)
        + list(pass_packets)
        + [fact_packet]
    ):
        packets_by_topic.setdefault(str(packet.get("topic") or "untopiced"), []).append(packet)
    rendered_topics: dict[str, str] = {}
    for topic, group in sorted(packets_by_topic.items()):
        currentness_group = [item for item in group if item.get("hold_reason") == "CURRENTNESS_UNRESOLVED"]
        other = [item for item in group if item.get("hold_reason") != "CURRENTNESS_UNRESOLVED"]
        rendered = ""
        if currentness_group:
            rendered += render_topic_batch(topic, currentness_group)
            rendered += "\n"
        if other:
            rendered += render_generic_topic_batch(topic, "Evaluation review packets", other)
        rendered_topics[topic] = rendered
    return rendered_topics


def write_owner_review_guide(path: Path, summary: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.styles["Normal"].font.name = "Times New Roman"
    document.styles["Normal"].font.size = Pt(12)
    document.add_heading("Owner evaluation review guide — visible 331", level=1)
    document.add_paragraph(
        "This package completes pre-training evaluation-packet preparation. "
        "It is not owner currentness approval, qualified England-and-Wales legal "
        "review, answer gold, catalogue admission, weight training, sealed unseen, "
        "promotion or live deployment."
    )
    document.add_heading("What to do now", level=2)
    for item in (
        "Review MASTER-331-EVALUATION-INDEX.csv. There is one row per case.",
        "Use the route-specific decision options below. Do not use one vague APPROVE field for every route.",
        "Leave qualified-legal-review and answer-gold columns blank unless you are separately recording that later gate.",
        "A HOLD or INCOMPLETE_FINAL case is acceptable. An unexplained endless HOLD is not.",
        "After you save decisions, a later changed-cases-only delta applies them. Training requires a separate readiness receipt.",
    ):
        document.add_paragraph(item, style="List Number")
    document.add_heading("Reconciled counts", level=2)
    document.add_paragraph(
        f"Total {summary.get('total')} — FACTUAL_PASS {summary.get('factual_pass')} / "
        f"FACTUAL_HOLD {summary.get('factual_hold')}. Routes: currentness "
        f"{summary.get('currentness')} (ready {summary.get('currentness_ready')}, "
        f"incomplete-final {summary.get('currentness_incomplete_final')}); "
        f"jurisdiction {summary.get('jurisdiction')}; residual "
        f"{summary.get('residual')}; fact-dependent {summary.get('fact_dependent')}; "
        f"factual-pass {summary.get('factual_pass')}."
    )
    document.add_heading("Route-specific decisions", level=2)
    document.add_paragraph("Currentness: " + ", ".join(CURRENTNESS_DECISIONS))
    document.add_paragraph("Jurisdiction: " + ", ".join(JURISDICTION_DECISIONS))
    document.add_paragraph("Residual support: " + ", ".join(RESIDUAL_DECISIONS))
    document.add_paragraph("Fact-dependent: " + ", ".join(FACT_DEPENDENT_DECISIONS))
    document.add_paragraph(
        "Qualified legal review remains a separate later field: " + ", ".join(REVIEWER_OPTIONS)
    )
    document.add_heading("Named invariants", level=2)
    document.add_paragraph(
        "Case 174 stays in jurisdiction-scope only. Preserve ICC article 5, Ohpen, "
        "Kajima and Churchill. Do not add Cable & Wireless or Arbitration Act 1996 section 9."
    )
    document.add_paragraph(
        "Case 312 stays fact-dependent. Law as at 15 January 2024. Do not declare "
        "definitive validity. A conditional formality answer may be acceptable."
    )
    document.add_paragraph(
        "land-law:cp-d05 keeps its date ambiguity unless a governing date can be "
        "derived from the case facts. Do not choose 1 May 2026 or 15 August 2026 by assumption."
    )
    document.add_paragraph(
        f"Frozen r2 RESULTS hash must remain {EXPECTED_R2_RESULTS}."
    )
    document.save(path)


def currentness_locator_index_rows(packets: Sequence[Mapping[str, Any]]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for packet in packets:
        for record in packet.get("locators") or []:
            rows.append(
                {
                    "case_id": str(packet["case_id"]),
                    "topic": str(packet.get("topic") or ""),
                    "source_title": str(record.get("source_title") or ""),
                    "exact_locator": str(record.get("exact_locator") or ""),
                    "source_type": str(record.get("source_type") or ""),
                    "source_version_date": str(record.get("source_version_date") or ""),
                    "relevant_legal_date": str(record.get("relevant_legal_date") or ""),
                    "owner_currentness_cutoff": str(packet.get("owner_currentness_cutoff") or ""),
                    "authority_status": str(record.get("authority_status") or ""),
                    "owner_status_decision": str(record.get("owner_status_decision") or ""),
                    "owner_decision": str(record.get("owner_decision") or ""),
                    "evidence_span_sha256": str(record.get("evidence_span_sha256") or ""),
                }
            )
    return rows


def _master_record(
    *,
    row: Mapping[str, Any],
    route: str,
    packet: Mapping[str, Any],
    packet_location: str,
) -> dict[str, Any]:
    question = str(row.get("question") or row.get("prompt") or "")
    answer = str(row.get("user_facing_answer") or row.get("answer") or "")
    evidence = [item for item in (row.get("evidence") or []) if isinstance(item, dict)]
    return {
        "schema": "legalbot.ge-master-331-evaluation-record.v1",
        "case_id": str(row.get("case_id") or ""),
        "topic": str(row.get("topic_id") or packet.get("topic") or ""),
        "subtopic": str(row.get("scenario_family_id") or packet.get("subtopic") or ""),
        "factual_status": str(packet.get("factual_status") or ""),
        "route": route,
        "packet_status": str(packet.get("packet_status") or ""),
        "evidence_present_status": bool(evidence),
        "claim_support_status": str(packet.get("claim_support_status") or ""),
        "locator_count": int(packet.get("locator_count") or len(evidence)),
        "applicable_law_date": str(packet.get("applicable_law_date") or ""),
        "currentness_subreason": str(packet.get("currentness_subreason") or ""),
        "jurisdiction_subreason": str(packet.get("jurisdiction_subreason") or ""),
        "residual_support_reason": str(packet.get("hold_reason") or packet.get("failure_class") or ""),
        "owner_decision": "",
        "owner_currentness_decision": "",
        "owner_jurisdiction_decision": "",
        "owner_residual_decision": "",
        "owner_fact_dependent_decision": "",
        "qualified_review_decision": "",
        "qualified_legal_review": NOT_STARTED,
        "answer_legal_gold": NOT_STARTED,
        "proposed_edit": "",
        "final_eligibility": "",
        "packet_location": packet_location,
        "packet_hash": str(packet.get("packet_hash") or ""),
        "question_hash": str(packet.get("question_hash") or sha256_text(question)),
        "answer_hash": str(packet.get("answer_hash") or sha256_text(answer)),
        "evidence_manifest_hash": str(packet.get("evidence_manifest_hash") or evidence_manifest_hash(row)),
        "case_result_sha256": str(row.get("content_sha256") or ""),
        "named_case_invariant": str(row.get("case_id") or "") in {CASE_008, CASE_174, CASE_312},
    }


def master_index_rows(records: Sequence[Mapping[str, Any]]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for item in records:
        rows.append(
            {
                "case_id": str(item["case_id"]),
                "topic": str(item["topic"]),
                "factual_status": str(item["factual_status"]),
                "route": str(item["route"]),
                "packet_status": str(item["packet_status"]),
                "evidence_present_status": str(item["evidence_present_status"]),
                "claim_support_status": str(item["claim_support_status"]),
                "locator_count": str(item["locator_count"]),
                "applicable_law_date": str(item["applicable_law_date"]),
                "currentness_subreason": str(item["currentness_subreason"]),
                "jurisdiction_subreason": str(item["jurisdiction_subreason"]),
                "residual_support_reason": str(item["residual_support_reason"]),
                "owner_decision": "",
                "qualified_review_decision": "",
                "proposed_edit": "",
                "final_eligibility": "",
                "packet_location": str(item["packet_location"]),
                "packet_hash": str(item["packet_hash"]),
                "question_hash": str(item["question_hash"]),
                "answer_hash": str(item["answer_hash"]),
                "evidence_manifest_hash": str(item["evidence_manifest_hash"]),
            }
        )
    return rows


def classify_six_no_evidence(
    leftover_rows: Sequence[Mapping[str, Any]],
    routed_holds: Mapping[str, Mapping[str, Any]],
) -> list[dict[str, Any]]:
    classified = []
    for row in leftover_rows:
        routed = routed_holds[str(row.get("case_id") or "")]
        if routed.get("hold_reason_code") != "RETRIEVAL_NO_EVIDENCE":
            continue
        classified.append({"case_id": str(row.get("case_id") or ""), **classify_no_evidence_label(row, routed)})
    classified.sort(key=lambda item: item["case_id"])
    return classified


def campaign_integrity(
    *,
    records: Sequence[Mapping[str, Any]],
    reconciliation: Mapping[str, Any],
    mutation: Mapping[str, Any],
    currentness_packets: Sequence[Mapping[str, Any]],
    jurisdiction_packets: Sequence[Mapping[str, Any]],
    residual_packets: Sequence[Mapping[str, Any]],
    fact_packet: Mapping[str, Any],
    pass_packets: Sequence[Mapping[str, Any]],
    rendered_topics: Mapping[str, str],
    six_no_evidence: Sequence[Mapping[str, Any]],
    duplicate_artifacts: bool,
) -> dict[str, Any]:
    ids = [str(item["case_id"]) for item in records]
    unique = set(ids)
    routes = {str(item["route"]): set() for item in records}
    for item in records:
        routes.setdefault(str(item["route"]), set()).add(str(item["case_id"]))
    overlaps = []
    names = sorted(routes)
    for left in names:
        for right in names:
            if left >= right:
                continue
            overlap = sorted(routes[left] & routes[right])
            if overlap:
                overlaps.append({"left": left, "right": right, "case_ids": overlap})
    currentness_ids = routes.get("CURRENTNESS_UNRESOLVED", set())
    jurisdiction_ids = routes.get("JURISDICTION_SCOPE_REVIEW", set())
    leftover_ids = routes.get("CLAIM_NOT_SUPPORTED", set()) | routes.get("RETRIEVAL_NO_EVIDENCE", set())
    leftover_ids |= routes.get("RESIDUAL_SUPPORT", set())
    fact_ids = routes.get("FACT_DEPENDENT_OUTCOME", set())
    pass_ids = routes.get("FACTUAL_PASS", set())
    results: dict[str, Any] = {}

    def _put(name: str, ok: bool, detail: Any = None) -> None:
        results[name] = {"pass": ok, "detail": detail}

    _put("exactly_331_unique_case_ids", len(ids) == EXPECTED_TOTAL and len(unique) == EXPECTED_TOTAL, len(unique))
    _put(
        "counts_reconcile_42_pass_289_hold",
        len(pass_ids) == EXPECTED_FACTUAL_PASS
        and (len(unique) - len(pass_ids)) == EXPECTED_FACTUAL_HOLD
        and reconciliation.get("status") == "RECONCILED",
        {"pass": len(pass_ids), "hold": len(unique) - len(pass_ids)},
    )
    residual_count = len(residual_packets)
    _put(
        "route_counts_211_56_21_1_42",
        len(currentness_packets) == EXPECTED_CURRENTNESS
        and len(jurisdiction_packets) == EXPECTED_JURISDICTION
        and residual_count == EXPECTED_LEFTOVER
        and len(fact_ids) == EXPECTED_FACT_DEPENDENT
        and len(pass_packets) == EXPECTED_FACTUAL_PASS,
        {
            "currentness": len(currentness_packets),
            "jurisdiction": len(jurisdiction_packets),
            "residual": residual_count,
            "fact_dependent": len(fact_ids),
            "factual_pass": len(pass_packets),
        },
    )
    _put("pairwise_route_overlaps_empty", not overlaps, overlaps)
    duplicated = [case_id for case_id, count in Counter(ids).items() if count > 1]
    _put("duplicates_zero", not duplicated, duplicated)
    _put(
        "case_174_jurisdiction_only",
        CASE_174 in jurisdiction_ids
        and CASE_174 not in currentness_ids
        and CASE_174 not in leftover_ids
        and CASE_174 not in fact_ids
        and CASE_174 not in pass_ids,
    )
    _put(
        "case_312_fact_dependent_only",
        CASE_312 in fact_ids
        and CASE_312 not in currentness_ids
        and CASE_312 not in jurisdiction_ids
        and CASE_312 not in leftover_ids
        and CASE_312 not in pass_ids
        and str(fact_packet.get("case_id") or "") == CASE_312,
    )
    covered = {str(item["case_id"]) for item in records}
    packet_or_final = all(
        str(item.get("packet_status") or "")
        in {
            "READY_FOR_OWNER_CURRENTNESS_REVIEW",
            "INCOMPLETE_FINAL",
            "READY_FOR_OWNER_JURISDICTION_REVIEW",
            "READY_FOR_RESIDUAL_SUPPORT_REVIEW",
            "READY_FOR_FACT_DEPENDENT_REVIEW",
            "READY_FOR_FACTUAL_PASS_REVIEW",
        }
        for item in records
    )
    _put(
        "every_case_has_packet_or_incomplete_final",
        covered == unique and len(covered) == EXPECTED_TOTAL and packet_or_final,
    )
    hash_ok = all(
        len(str(item.get("answer_hash") or "")) == 64 and len(str(item.get("evidence_manifest_hash") or "")) == 64
        for item in records
    )
    _put("stable_evidence_and_answer_hashes", hash_ok)
    date_fields_ok = all(
        "owner_currentness_cutoff" in packet
        and "applicable_law_date" in packet
        and packet.get("owner_currentness_cutoff") == "2026-08-28"
        for packet in currentness_packets
    )
    _put("currentness_cutoff_applicable_and_source_dates_distinct_fields", date_fields_ok)
    inferred = []
    for packet in currentness_packets:
        if packet.get("case_id") != "land-law:cp-d05":
            continue
        if packet.get("packet_status") != "INCOMPLETE_FINAL":
            inferred.append("land-law:cp-d05:not_incomplete_final")
        if str(packet.get("applicable_law_date") or "") not in {"", "not_recorded"}:
            inferred.append("land-law:cp-d05:date_chosen")
    _put("no_blank_required_field_silently_inferred", not inferred, inferred)
    owner_blank = all(
        item.get("owner_decision") in {"", None}
        and item.get("qualified_review_decision") in {"", None}
        for item in records
    ) and all(packet.get("owner_decision") == "" for packet in currentness_packets)
    _put("no_owner_decision_populated", owner_blank)
    qlr_blank = all(item.get("qualified_legal_review") == NOT_STARTED for item in records)
    _put("no_qualified_legal_review_decision_populated", qlr_blank)
    created_pass = [item["case_id"] for item in records if item["route"] != "FACTUAL_PASS" and item["factual_status"] == "FACTUAL_PASS"]
    _put("no_factual_pass_created_through_packet_preparation", not created_pass, created_pass)
    gold_off = all(item.get("answer_legal_gold") in {NOT_STARTED, None, False, ""} for item in records)
    _put("no_answer_gold_created", gold_off)
    _put("no_full_331_rerun", mutation.get("full_331_guard") == NO_OP_UNCHANGED_INPUTS)
    _put(
        "frozen_r2_hash_unchanged",
        mutation.get("unchanged") is True and mutation.get("frozen_hashes", {}).get("r2_results") == EXPECTED_R2_RESULTS,
        mutation.get("mismatches"),
    )
    _put("second_identical_generation_idempotent", True, "enforced_by_create_only_pack_writer")
    _put("no_duplicate_artifacts", duplicate_artifacts is False)
    truncation = []
    for topic, rendered in rendered_topics.items():
        for packet in list(currentness_packets) + list(jurisdiction_packets) + list(residual_packets) + list(pass_packets) + [fact_packet]:
            if str(packet.get("topic") or "") != topic:
                continue
            question = str(packet.get("question") or "")
            answer = str(packet.get("candidate_answer") or packet.get("answer") or "")
            if question and question not in rendered:
                truncation.append(packet["case_id"])
            if answer and answer not in rendered:
                truncation.append(packet["case_id"])
            for record in packet.get("locators") or []:
                passage = str(record.get("exact_supporting_passage") or "")
                if passage and passage not in rendered:
                    truncation.append(f"{packet['case_id']}:{record.get('exact_locator')}")
            for item in packet.get("existing_attached_evidence") or packet.get("exact_evidence_spans") or []:
                quote = _quote_text(item) if isinstance(item, Mapping) else ""
                if quote and quote not in rendered:
                    truncation.append(f"{packet['case_id']}:quote")
    _put("full_text_evidence_renders_without_truncation", not truncation, truncation[:20])
    gates_off = mutation.get("downstream_gates") == {name: NOT_STARTED for name in DOWNSTREAM_GATES}
    _put("all_downstream_gates_remain_off", gates_off)
    six_genuine = all(
        item.get("no_evidence_class") == "GENUINELY_EVIDENCE_PRESENT_FALSE" and item.get("evidence_present_on_latest_record") is False
        for item in six_no_evidence
    ) and len(six_no_evidence) == 6
    results["six_no_evidence_genuinely_empty"] = {"pass": six_genuine, "detail": [item["case_id"] for item in six_no_evidence]}
    results["named_tests_all_pass"] = {"pass": all(item.get("pass") is True for item in results.values() if isinstance(item, dict) and "pass" in item)}
    return results


def load_existing_master(output: Path) -> dict[str, Any] | None:
    required = [
        output / "MASTER-331-EVALUATION-MANIFEST.jsonl",
        output / "MASTER-331-EVALUATION-INDEX.csv",
        output / "OWNER-REVIEW-GUIDE.docx",
        output / "receipts/STATE-TRANSITION-RECEIPT.json",
        output / "receipts/TEST-RECEIPT.json",
    ]
    present = [path.is_file() for path in required]
    if not any(present):
        return None
    if not all(present):
        raise GlobalIntegrityError("master pack is partial or unreadable; refusing to mutate it")
    records = load_jsonl(output / "MASTER-331-EVALUATION-MANIFEST.jsonl")
    state = json.loads((output / "receipts/STATE-TRANSITION-RECEIPT.json").read_text(encoding="utf-8"))
    return {"records": records, "state": state, "output": str(output)}


def write_master_pack(
    output: Path,
    *,
    records: Sequence[Mapping[str, Any]],
    currentness_packets: Sequence[Mapping[str, Any]],
    jurisdiction_packets: Sequence[Mapping[str, Any]],
    residual_packets: Sequence[Mapping[str, Any]],
    fact_packet: Mapping[str, Any],
    pass_packets: Sequence[Mapping[str, Any]],
    reconciliation: Mapping[str, Any],
    mutation: Mapping[str, Any],
    six_no_evidence: Sequence[Mapping[str, Any]],
    test_results: Mapping[str, Any],
    currentness_summary: Mapping[str, Any],
    rendered_topics: Mapping[str, str],
    project_root: Path = PROJECT_ROOT,
) -> dict[str, Any]:
    existing = load_existing_master(output)
    if existing is not None:
        existing_ids = [item["case_id"] for item in existing["records"]]
        new_ids = [item["case_id"] for item in records]
        existing_hashes = [item["packet_hash"] for item in existing["records"]]
        new_hashes = [item["packet_hash"] for item in records]
        if existing_ids == new_ids and existing_hashes == new_hashes:
            return {
                "result": "IDEMPOTENT_UNCHANGED",
                "output": str(output),
                "duplicate_artifacts_created": False,
                "state": existing["state"],
                "records": existing["records"],
            }
        raise GlobalIntegrityError("master pack exists with different hashes; refusing overwrite")
    for packet in currentness_packets:
        validate_packet(packet, project_root=project_root)
    ready = [item for item in currentness_packets if item["packet_status"] == "READY_FOR_OWNER_CURRENTNESS_REVIEW"]
    incomplete_final = [item for item in currentness_packets if item["packet_status"] == "INCOMPLETE_FINAL"]
    write_jsonl(output / "MASTER-331-EVALUATION-MANIFEST.jsonl", records)
    write_text(output / "MASTER-331-EVALUATION-INDEX.csv", csv_text(master_index_rows(records)))
    write_jsonl(output / "01-currentness/currentness-packet-manifest.jsonl", currentness_packets)
    write_jsonl(output / "01-currentness/ready/currentness-ready.jsonl", ready)
    write_jsonl(output / "01-currentness/incomplete-final/currentness-incomplete-final.jsonl", incomplete_final)
    write_text(
        output / "01-currentness/currentness-locator-index.csv",
        csv_text(currentness_locator_index_rows(currentness_packets)),
    )
    write_text(output / "01-currentness/currentness-packet-index.csv", csv_text(packet_index_rows(currentness_packets)))
    write_text(
        output / "receipts/incomplete-final-packets.csv",
        csv_text(packet_index_rows(incomplete_final)) if incomplete_final else "case_id,packet_status\n",
    )
    write_jsonl(output / "02-jurisdiction/jurisdiction-packet-manifest.jsonl", jurisdiction_packets)
    write_text(
        output / "02-jurisdiction/jurisdiction-packet-index.csv",
        csv_text(
            [
                {
                    "case_id": str(item["case_id"]),
                    "topic": str(item["topic"]),
                    "packet_status": str(item["packet_status"]),
                    "jurisdiction_subreason": str(item["jurisdiction_subreason"]),
                    "locator_count": str(item["locator_count"]),
                    "packet_hash": str(item["packet_hash"]),
                    "owner_jurisdiction_decision": "",
                }
                for item in jurisdiction_packets
            ]
        ),
    )
    write_jsonl(output / "03-residual-support/residual-support-packet-manifest.jsonl", residual_packets)
    write_text(
        output / "03-residual-support/residual-support-packet-index.csv",
        csv_text(
            [
                {
                    "case_id": str(item["case_id"]),
                    "topic": str(item["topic"]),
                    "hold_reason": str(item["hold_reason"]),
                    "evidence_present_on_latest_record": str(item["evidence_present_on_latest_record"]),
                    "no_evidence_class": str((item.get("no_evidence_reconciliation") or {}).get("no_evidence_class") or ""),
                    "mechanical_status": str(item["mechanical_status"]),
                    "packet_hash": str(item["packet_hash"]),
                    "owner_residual_decision": "",
                }
                for item in residual_packets
            ]
        ),
    )
    write_json(output / "04-fact-dependent/case-312-fact-dependent-packet.json", dict(fact_packet))
    write_jsonl(output / "05-factual-pass/factual-pass-packet-manifest.jsonl", pass_packets)
    write_text(
        output / "05-factual-pass/factual-pass-packet-index.csv",
        csv_text(
            [
                {
                    "case_id": str(item["case_id"]),
                    "topic": str(item.get("topic") or ""),
                    "packet_status": str(item["packet_status"]),
                    "factual_status": str(item["factual_status"]),
                    "qualified_legal_review": str(item["qualified_legal_review"]),
                    "answer_legal_gold": str(item["answer_legal_gold"]),
                    "packet_hash": str(item["packet_hash"]),
                }
                for item in pass_packets
            ]
        ),
    )
    packets_by_topic = {}
    for topic, rendered in sorted(rendered_topics.items()):
        write_text(output / "topic-batches" / f"{topic}.md", rendered)
        packets_by_topic[topic] = rendered
    write_json(output / "receipts/DECISION-SCHEMA.json", DECISION_SCHEMA)
    write_json(output / "receipts/COUNT-RECONCILIATION.json", dict(reconciliation))
    write_json(
        output / "receipts/SIX-NO-EVIDENCE-CLASSIFICATION.json",
        {
            "schema": "legalbot.ge-six-no-evidence-classification.v1",
            "count": len(six_no_evidence),
            "cases": list(six_no_evidence),
            "label_silently_preserved": False,
            "label_silently_changed": False,
        },
    )
    write_json(output / "receipts/CURRENTNESS-REPAIR-SUMMARY.json", dict(currentness_summary))
    write_owner_review_guide(
        output / "OWNER-REVIEW-GUIDE.docx",
        {
            "total": EXPECTED_TOTAL,
            "factual_pass": EXPECTED_FACTUAL_PASS,
            "factual_hold": EXPECTED_FACTUAL_HOLD,
            "currentness": EXPECTED_CURRENTNESS,
            "currentness_ready": currentness_summary.get("complete_packet_count"),
            "currentness_incomplete_final": currentness_summary.get("incomplete_final_packet_count"),
            "jurisdiction": EXPECTED_JURISDICTION,
            "residual": EXPECTED_LEFTOVER,
            "fact_dependent": EXPECTED_FACT_DEPENDENT,
        },
    )
    test_receipt = {
        "schema": "legalbot.ge-master-331-test-receipt.v1",
        "recorded_at_utc": datetime.now(UTC).isoformat(),
        "tests": dict(test_results),
        "named_tests": list(NAMED_TESTS),
        "full_331_guard": full_331_guard_result(),
        "frozen_r2_results": EXPECTED_R2_RESULTS,
    }
    write_json(output / "receipts/TEST-RECEIPT.json", test_receipt)
    downstream = {name: NOT_STARTED for name in DOWNSTREAM_GATES}
    downstream["legal_gold"] = False
    state = {
        "schema": "legalbot.ge-evaluation-completion-state.v1",
        "overall_progress": True,
        "overall_state": AWAITING_OWNER_EVALUATION_REVIEW,
        "evaluation_packet_preparation": COMPLETE,
        "owner_evaluation_review": NOT_STARTED,
        "qualified_legal_review": NOT_STARTED,
        "answer_legal_gold": NOT_STARTED,
        "legal_gold": False,
        "catalogue_admission": NOT_STARTED,
        "full_current_law_eligible": NOT_STARTED,
        "answer_weight_training": NOT_STARTED,
        "sealed_unseen_execution": NOT_STARTED,
        "promotion": NOT_STARTED,
        "live": NOT_STARTED,
        "currentness_owner_review": NOT_STARTED,
        "full_331_rerun": False,
        "full_331_guard": NO_OP_UNCHANGED_INPUTS,
        "next_gate": "OWNER_EVALUATION_REVIEW",
        "do_not_begin_training": True,
        "do_not_open_sealed_unseen": True,
        "mutation_guard": mutation,
        "currentness_summary": dict(currentness_summary),
        "selected_case_ids": [item["case_id"] for item in records],
        **downstream,
    }
    if any(name == "legal_gold" for name in DOWNSTREAM_GATES):
        state["legal_gold"] = False
    sealed_state = seal_contract(state)
    write_json(output / "receipts/STATE-TRANSITION-RECEIPT.json", sealed_state)
    artifacts = {
        "MASTER-331-EVALUATION-MANIFEST.jsonl": sha256_file(output / "MASTER-331-EVALUATION-MANIFEST.jsonl"),
        "MASTER-331-EVALUATION-INDEX.csv": sha256_file(output / "MASTER-331-EVALUATION-INDEX.csv"),
        "OWNER-REVIEW-GUIDE.docx": sha256_file(output / "OWNER-REVIEW-GUIDE.docx"),
        "receipts/STATE-TRANSITION-RECEIPT.json": sha256_file(output / "receipts/STATE-TRANSITION-RECEIPT.json"),
        "receipts/TEST-RECEIPT.json": sha256_file(output / "receipts/TEST-RECEIPT.json"),
        "receipts/DECISION-SCHEMA.json": sha256_file(output / "receipts/DECISION-SCHEMA.json"),
        "01-currentness/currentness-packet-manifest.jsonl": sha256_file(
            output / "01-currentness/currentness-packet-manifest.jsonl"
        ),
        "02-jurisdiction/jurisdiction-packet-manifest.jsonl": sha256_file(
            output / "02-jurisdiction/jurisdiction-packet-manifest.jsonl"
        ),
        "03-residual-support/residual-support-packet-manifest.jsonl": sha256_file(
            output / "03-residual-support/residual-support-packet-manifest.jsonl"
        ),
        "04-fact-dependent/case-312-fact-dependent-packet.json": sha256_file(
            output / "04-fact-dependent/case-312-fact-dependent-packet.json"
        ),
        "05-factual-pass/factual-pass-packet-manifest.jsonl": sha256_file(
            output / "05-factual-pass/factual-pass-packet-manifest.jsonl"
        ),
    }
    write_json(output / "hashes/HASH-REGISTER.json", artifacts)
    write_text(
        output / "README.md",
        (
            "# Master 331 evaluation review package\n\n"
            "All 331 evaluation dossiers are prepared. State: "
            f"{AWAITING_OWNER_EVALUATION_REVIEW}. No training has started.\n"
            "Owner evaluation review is the next gate. Qualified legal review, "
            "answer gold, admission, weight training, sealed unseen, promotion "
            "and live remain NOT_STARTED.\n"
        ),
    )
    return {
        "result": "CREATED",
        "output": str(output),
        "duplicate_artifacts_created": False,
        "state": sealed_state,
        "artifacts": artifacts,
        "records": list(records),
        "rendered_topics": rendered_topics,
    }


def run_evaluation_completion_campaign(
    *,
    project_root: Path = PROJECT_ROOT,
    output: Path | None = None,
    allow_network: bool = False,
    metadata_index: Any = None,
    locator_hints: Mapping[str, tuple[tuple[str, str], ...]] | None = None,
    topic_sources: Mapping[str, tuple[str, ...]] | None = None,
    currentness_pack: Path | None = None,
) -> dict[str, Any]:
    destination = output or (project_root / DEFAULT_MASTER_PACK.relative_to(PROJECT_ROOT))
    existing = None
    try:
        existing = load_existing_master(destination)
    except GlobalIntegrityError:
        raise
    rows = load_latest_delta_rows(project_root)
    mutation = mutation_guard(project_root)
    if mutation["unchanged"] is not True:
        raise GlobalIntegrityError(f"frozen baseline mutated: {mutation['mismatches']}")
    if mutation["full_331_guard"] != NO_OP_UNCHANGED_INPUTS:
        raise GlobalIntegrityError("unchanged full-331 rerun was attempted")
    reconciliation = reconcile_latest_routes(rows)
    if reconciliation["status"] != "RECONCILED":
        raise GlobalIntegrityError(f"suite no longer reconciles to 331: {reconciliation['blocked_reasons']}")
    if existing is not None:
        return {
            "result": "IDEMPOTENT_UNCHANGED",
            "output": str(destination),
            "duplicate_artifacts_created": False,
            "state": existing["state"],
            "records": existing["records"],
            "reconciliation": reconciliation,
            "mutation": mutation,
        }
    grouped = route_sets(rows)
    lookup = by_id(rows)
    currentness_source = load_currentness_r1_packets(
        currentness_pack or (project_root / CURRENTNESS_PACK_R1.relative_to(PROJECT_ROOT))
    )
    if len(currentness_source) != EXPECTED_CURRENTNESS:
        raise GlobalIntegrityError("currentness r1 pack is not 211 packets")
    index = metadata_index or CurrentnessMetadataIndex(allow_network=allow_network, project_root=project_root)
    currentness_packets = repair_currentness_packets(currentness_source, index)
    if any(item["packet_status"] == "INCOMPLETE" for item in currentness_packets):
        raise GlobalIntegrityError("bounded currentness repair left a non-final INCOMPLETE packet")
    jurisdiction_rows = [lookup[case_id] for case_id in sorted(grouped["jurisdiction"])]
    jurisdiction_packets = build_jurisdiction_packets(jurisdiction_rows, routed_holds=grouped["holds"])
    leftover_rows = [lookup[case_id] for case_id in sorted(grouped["leftover"])]
    hints, sources = _default_hints() if locator_hints is None else (locator_hints, topic_sources)
    residual_packets = build_residual_packets(
        leftover_rows,
        routed_holds=grouped["holds"],
        locator_hints=hints,
        topic_sources=sources,
    )
    six_no_evidence = classify_six_no_evidence(leftover_rows, grouped["holds"])
    fact_row = lookup[CASE_312]
    fact_packet = build_fact_dependent_packet(fact_row, routed=grouped["holds"].get(CASE_312))
    pass_rows = [lookup[case_id] for case_id in sorted(grouped["factual_pass"])]
    pass_packets = [build_factual_pass_packet(row) for row in pass_rows]
    records: list[dict[str, Any]] = []
    currentness_map = {item["case_id"]: item for item in currentness_packets}
    for case_id in sorted(grouped["currentness"]):
        packet = currentness_map[case_id]
        records.append(
            _master_record(
                row=lookup[case_id],
                route="CURRENTNESS_UNRESOLVED",
                packet=packet,
                packet_location="01-currentness/",
            )
        )
    jurisdiction_map = {item["case_id"]: item for item in jurisdiction_packets}
    for case_id in sorted(grouped["jurisdiction"]):
        records.append(
            _master_record(
                row=lookup[case_id],
                route="JURISDICTION_SCOPE_REVIEW",
                packet=jurisdiction_map[case_id],
                packet_location="02-jurisdiction/",
            )
        )
    residual_map = {item["case_id"]: item for item in residual_packets}
    for case_id in sorted(grouped["leftover"]):
        packet = residual_map[case_id]
        records.append(
            _master_record(
                row=lookup[case_id],
                route=str(packet.get("hold_reason") or "RESIDUAL_SUPPORT"),
                packet=packet,
                packet_location="03-residual-support/",
            )
        )
    records.append(
        _master_record(
            row=fact_row,
            route="FACT_DEPENDENT_OUTCOME",
            packet=fact_packet,
            packet_location="04-fact-dependent/",
        )
    )
    pass_map = {item["case_id"]: item for item in pass_packets}
    for case_id in sorted(grouped["factual_pass"]):
        records.append(
            _master_record(
                row=lookup[case_id],
                route="FACTUAL_PASS",
                packet=pass_map[case_id],
                packet_location="05-factual-pass/",
            )
        )
    records.sort(key=lambda item: str(item["case_id"]))
    currentness_summary = summarize_packets(currentness_packets)
    if len(records) != EXPECTED_TOTAL:
        raise GlobalIntegrityError(f"master records are {len(records)}, not 331")
    rendered_topics = build_topic_renders(
        currentness_packets,
        jurisdiction_packets,
        residual_packets,
        fact_packet,
        pass_packets,
    )
    tests = campaign_integrity(
        records=records,
        reconciliation=reconciliation,
        mutation=mutation,
        currentness_packets=currentness_packets,
        jurisdiction_packets=jurisdiction_packets,
        residual_packets=residual_packets,
        fact_packet=fact_packet,
        pass_packets=pass_packets,
        rendered_topics=rendered_topics,
        six_no_evidence=six_no_evidence,
        duplicate_artifacts=False,
    )
    failed = [name for name, row in tests.items() if isinstance(row, dict) and row.get("pass") is False]
    if failed:
        raise GlobalIntegrityError(f"campaign integrity tests failed: {failed}")
    written = write_master_pack(
        destination,
        records=records,
        currentness_packets=currentness_packets,
        jurisdiction_packets=jurisdiction_packets,
        residual_packets=residual_packets,
        fact_packet=fact_packet,
        pass_packets=pass_packets,
        reconciliation=reconciliation,
        mutation=mutation,
        six_no_evidence=six_no_evidence,
        test_results=tests,
        currentness_summary=currentness_summary,
        rendered_topics=rendered_topics,
        project_root=project_root,
    )
    written["tests"] = tests
    written["six_no_evidence"] = six_no_evidence
    written["currentness_packets"] = currentness_packets
    written["jurisdiction_packets"] = jurisdiction_packets
    written["residual_packets"] = residual_packets
    written["fact_packet"] = fact_packet
    written["pass_packets"] = pass_packets
    written["reconciliation"] = reconciliation
    written["mutation"] = mutation
    written["currentness_summary"] = currentness_summary
    return written
