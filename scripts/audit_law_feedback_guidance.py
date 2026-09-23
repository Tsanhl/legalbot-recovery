#!/usr/bin/env python3
"""Inventory Law feedback without copying student work into the shared guide.

The output is source custody and extraction coverage, not legal authority or
automatic permission to admit a feedback-derived runtime rule.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

CRITERIA = {
    "analysis": ("analysis", "analytical", "critical", "reasoning"),
    "application": ("application", "apply", "facts", "scenario"),
    "authority": ("authority", "authorities", "cases", "statute"),
    "pinpoint": ("pinpoint", "oscola", "footnote", "reference"),
    "structure": ("structure", "headings", "organisation", "conclusion"),
    "counterargument": ("counterargument", "compare", "contrast", "alternative"),
    "clarity": ("clarity", "concise", "grammar", "proofread"),
}


def _candidate(root: Path, path: Path) -> bool:
    relative = str(path.relative_to(root)).casefold()
    return "feedback" in relative or "assessment criteria" in relative or "marking criteria" in relative


def _extract(path: Path) -> tuple[str, int, int]:
    suffix = path.suffix.casefold()
    if suffix == ".pdf":
        reader = PdfReader(path)
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages), len(pages), sum(len(page.strip()) < 70 for page in pages)
    if suffix == ".docx":
        document = Document(path)
        text = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            text.extend(cell.text for row in table.rows for cell in row.cells)
        return "\n".join(text), 0, 0
    if suffix == ".pptx":
        presentation = Presentation(path)
        text = [shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame]
        return "\n".join(text), len(presentation.slides), 0
    raise ValueError("unsupported type")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--law-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    root = args.law_root.resolve()
    if not root.is_dir() or args.output.exists():
        parser.error("Law directory must exist and output must be new")
    records = []
    totals: Counter[str] = Counter()
    for path in sorted(root.rglob("*")):
        if not path.is_file() or not _candidate(root, path):
            continue
        if path.suffix.casefold() == ".zip":
            with ZipFile(path) as archive:
                count = sum(not item.is_dir() for item in archive.infolist())
            records.append({"path": str(path.relative_to(root)), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                            "format": "zip", "archive_members": count, "status": "archive_inventory_only"})
            totals["archive_inventory_only"] += 1
            continue
        if path.suffix.casefold() not in {".pdf", ".docx", ".pptx"}:
            continue
        try:
            text, pages, short_pages = _extract(path)
            folded = " ".join(text.casefold().split())
            signals = {name: sum(len(re.findall(rf"\b{re.escape(term)}\b", folded)) for term in terms)
                       for name, terms in CRITERIA.items()}
            status = "text_partial_ocr_needed" if short_pages else "text_extracted"
            records.append({
                "path": str(path.relative_to(root)), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                "format": path.suffix.casefold().lstrip("."), "page_or_slide_count": pages,
                "short_or_image_only_pages": short_pages, "extracted_character_count": len(text),
                "criterion_signals": signals, "status": status,
            })
            totals[status] += 1
            totals["short_or_image_only_pages"] += short_pages
        except Exception as exc:
            records.append({"path": str(path.relative_to(root)), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                            "format": path.suffix.casefold().lstrip("."), "status": "extraction_failed",
                            "error_type": type(exc).__name__})
            totals["extraction_failed"] += 1
    output = {
        "schema": "legalbot.law-feedback-extraction-audit.v1",
        "recorded_at": datetime.now(UTC).isoformat(),
        "law_root": str(root), "document_count": len(records), "totals": dict(totals),
        "records": records, "raw_student_work_retained": False,
        "runtime_guide_admission": False,
        "limit": "Filename/path discovery and extracted text only; image-only pages require OCR and reviewed generalisation.",
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("x", encoding="utf-8") as handle:
        json.dump(output, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
    os.chmod(args.output, 0o600)
    print(json.dumps({"document_count": len(records), "totals": dict(totals)}, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
